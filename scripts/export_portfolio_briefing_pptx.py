#!/usr/bin/env python3
"""Export the 0907 HTML briefing to a 16:9 leadership PPTX.

Content is unchanged: each HTML slide is captured as-is. UI chrome that is
not briefing content (language toggle, keyboard hint) is hidden. Scrollable
cost trees and the equipment table are expanded and split across extra 16:9
slides so a projector does not clip numbers.
"""
from __future__ import annotations

import http.server
import os
import socketserver
import threading
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches
from playwright.sync_api import sync_playwright

ROOT = Path("/workspace")
HTML_REL = "汇报/UK-PDF-Portfolio/UK_PDF_Portfolio_Briefing_2026-09-07.html"
OUT = ROOT / "汇报/UK-PDF-Portfolio/UK_PDF_Portfolio_Briefing_2026-09-07.pptx"
SHOT_DIR = Path("/tmp/portfolio-ppt-shots")
PORT = 0  # ephemeral; actual port taken from the bound server
VW, VH = 1920, 1080
DSF = 2
PX_W, PX_H = VW * DSF, VH * DSF
BG = (244, 246, 248)  # --bg #f4f6f8


PREPARE_JS = """
() => {
  const style = document.createElement('style');
  style.textContent = `
    body, button, input, table, .slide {
      font-family: "WenQuanYi Micro Hei", "Segoe UI", sans-serif !important;
    }
    .lang-switch { display: none !important; }
    #navHint { display: none !important; }
    .equip-filters { display: none !important; }
    .slide { animation: none !important; }
    .cost-scroll, .equip-table-wrap { scrollbar-width: none; }
    .cost-scroll::-webkit-scrollbar, .equip-table-wrap::-webkit-scrollbar { display: none; }
  `;
  document.head.appendChild(style);
  document.getElementById('navHint')?.remove();
  document.documentElement.classList.remove('en');
}
"""


def serve() -> socketserver.TCPServer:
    os.chdir(ROOT)

    class Handler(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *_a):
            pass

    class ReuseServer(socketserver.TCPServer):
        allow_reuse_address = True

    httpd = ReuseServer(("127.0.0.1", PORT), Handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def wait_ready(page) -> None:
    page.wait_for_function("() => typeof Chart !== 'undefined' && typeof show === 'function'")
    page.wait_for_timeout(200)


def goto_slide(page, index: int) -> None:
    page.evaluate("(n) => { show(n); }", index)
    page.wait_for_timeout(250)
    # Chart.js pages need a paint after buildCharts()
    page.wait_for_function(
        """() => {
          const s = document.querySelector('.slide.active');
          if (!s || !s.dataset.charts) return true;
          const cvs = s.querySelectorAll('canvas');
          if (!cvs.length) return true;
          return [...cvs].every(c => c.width > 20 && c.height > 20);
        }"""
    )
    page.wait_for_timeout(400)


def save_tile(im: Image.Image, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if im.mode != "RGB":
        im = im.convert("RGB")
    im.save(dest, "JPEG", quality=93, optimize=True, progressive=True)


def slice_to_16x9(im: Image.Image, stem: str) -> list[Path]:
    """Split a tall capture into 16:9 JPEG tiles with a small overlap."""
    if im.width != PX_W:
        im = im.resize((PX_W, int(im.height * PX_W / im.width)), Image.Resampling.LANCZOS)
    if im.height <= PX_H + 8:
        out = SHOT_DIR / f"{stem}.jpg"
        canvas = Image.new("RGB", (PX_W, PX_H), BG)
        canvas.paste(im, (0, 0))
        save_tile(canvas, out)
        return [out]

    overlap = 120
    tiles = []
    y0 = 0
    i = 1
    while y0 < im.height:
        y1 = min(y0 + PX_H, im.height)
        if y1 - y0 < PX_H and y0 > 0:
            y0 = max(0, im.height - PX_H)
            y1 = im.height
        canvas = Image.new("RGB", (PX_W, PX_H), BG)
        part = im.crop((0, y0, PX_W, y1))
        canvas.paste(part, (0, 0))
        out = SHOT_DIR / f"{stem}_{i:02d}.jpg"
        save_tile(canvas, out)
        tiles.append(out)
        if y1 >= im.height:
            break
        y0 = y1 - overlap
        i += 1
        if i > 8:
            break
    return tiles


def scroll_info(page) -> dict:
    return page.evaluate(
        """() => {
          const slide = document.querySelector('.slide.active');
          const el = slide && slide.querySelector('.cost-scroll, .equip-table-wrap');
          if (!el) return {tall: false, scrollHeight: 0, clientHeight: 0};
          return {
            tall: el.scrollHeight > el.clientHeight + 12,
            scrollHeight: el.scrollHeight,
            clientHeight: el.clientHeight,
          };
        }"""
    )


def set_scroll(page, y: int) -> None:
    page.evaluate(
        """(y) => {
          const el = document.querySelector('.slide.active .cost-scroll, .slide.active .equip-table-wrap');
          if (el) el.scrollTop = y;
        }""",
        y,
    )


def capture_slide(page, index: int) -> list[Path]:
    goto_slide(page, index)
    stem = f"p{index + 1:02d}"
    info = scroll_info(page)
    if not info["tall"]:
        png = page.screenshot(type="png", full_page=False)
        im = Image.open(BytesIO(png))
        return slice_to_16x9(im, stem)

    # Keep title + footer on every tile; page the inner scroll pane.
    overlap = 48
    step = max(80, int(info["clientHeight"] - overlap))
    max_y = int(info["scrollHeight"] - info["clientHeight"])
    tiles: list[Path] = []
    y = 0
    i = 1
    while True:
        set_scroll(page, y)
        page.wait_for_timeout(120)
        png = page.screenshot(type="png", full_page=False)
        im = Image.open(BytesIO(png))
        out = SHOT_DIR / f"{stem}_{i:02d}.jpg"
        save_tile(im.convert("RGB"), out)
        tiles.append(out)
        if y >= max_y:
            break
        nxt = y + step
        y = max_y if nxt >= max_y else nxt
        i += 1
        if i > 8:
            break
    set_scroll(page, 0)
    return tiles


def add_image_slide(prs: Presentation, img: Path) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(img), Emu(0), Emu(0), prs.slide_width, prs.slide_height)


def build(limit: int | None = None) -> None:
    SHOT_DIR.mkdir(parents=True, exist_ok=True)
    httpd = serve()
    port = httpd.server_address[1]
    url = f"http://127.0.0.1:{port}/{HTML_REL}"
    paths: list[Path] = []
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(
                channel="chrome",
                headless=True,
                args=[
                    "--no-sandbox",
                    "--disable-gpu",
                    "--disable-dev-shm-usage",
                    "--disable-background-networking",
                    "--disable-sync",
                    "--disable-extensions",
                    "--hide-scrollbars",
                ],
            )
            context = browser.new_context(
                viewport={"width": VW, "height": VH},
                device_scale_factor=DSF,
                color_scheme="light",
            )
            page = context.new_page()
            page.goto(url, wait_until="networkidle", timeout=60000)
            page.evaluate(PREPARE_JS)
            wait_ready(page)
            n = page.evaluate("() => document.querySelectorAll('.slide').length")
            last = n if limit is None else min(n, limit)
            print(f"slides in HTML: {n}; capturing 1..{last}", flush=True)
            for i in range(last):
                print(f"capture {i + 1}/{last}", flush=True)
                tiles = capture_slide(page, i)
                paths.extend(tiles)
                for t in tiles:
                    print(f"  -> {t.name} {t.stat().st_size}", flush=True)
            context.close()
            browser.close()
    finally:
        httpd.shutdown()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for img in paths:
        add_image_slide(prs, img)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print("wrote", OUT, "slides", len(prs.slides), "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    import sys

    lim = int(sys.argv[1]) if len(sys.argv) > 1 else None
    build(lim)
