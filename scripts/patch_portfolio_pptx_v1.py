#!/usr/bin/env python3
"""Patch the uploaded 9.7 v1 PPT: native Gantt pages + split the overflowing CRD table.

Preserves the user's wording. Writes back to v1 and the canonical 0907 file.
"""
from __future__ import annotations

import re
import sys
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_portfolio_pptx_zh import (  # noqa: E402
    CW,
    HIPO_C,
    ML,
    MUTED,
    NAVY,
    ROW_ALT,
    TEXT,
    WHITE,
    fill_cell,
)
from pptx_gantt import GANTT_SPECS, draw_gantt  # noqa: E402

ROOT = Path("/workspace")
V1 = ROOT / "汇报/UK-PDF-Portfolio/UK_PDF_Portfolio_Briefing_2026-09-07-v1.pptx"
OUT = ROOT / "汇报/UK-PDF-Portfolio/UK_PDF_Portfolio_Briefing_2026-09-07.pptx"

GANTT_SLIDES = {
    8: "ext",
    15: "hplc",
    22: "c1",
    33: "hipo",  # 0-based; becomes 34 after CRD insert — applied before insert
}


def _emu_in(v) -> float:
    return int(v) / 914400


def _remove(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def replace_gantt(slide, spec: dict) -> None:
    pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    cards = []
    for sh in slide.shapes:
        if sh.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        if _emu_in(sh.width) > 10 and _emu_in(sh.height) > 3.5 and 1.0 < _emu_in(sh.top) < 2.5:
            cards.append(sh)
    if not cards:
        raise RuntimeError("gantt card not found")
    card = max(cards, key=lambda sh: int(sh.width) * int(sh.height))
    box = (_emu_in(card.left), _emu_in(card.top), _emu_in(card.width), _emu_in(card.height))
    for sh in pics + cards:
        _remove(sh)
    draw_gantt(slide, *box, spec)


def _clone_slide(prs, index: int):
    source = prs.slides[index]
    dest = prs.slides.add_slide(prs.slide_layouts[6])
    for sh in list(dest.shapes):
        _remove(sh)
    src_tree = source.shapes._spTree
    dst_tree = dest.shapes._spTree
    for child in list(src_tree):
        tag = child.tag
        if tag.endswith("}nvGrpSpPr") or tag.endswith("}grpSpPr"):
            continue
        dst_tree.append(deepcopy(child))
    return dest


def _move_slide(prs, old_index: int, new_index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst
    entries = list(sld_id_lst)
    el = entries[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def _set_run_text(shape, new: str) -> None:
    tf = shape.text_frame
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = new
        for extra in p.runs[1:]:
            extra.text = ""
    else:
        p.text = new


def _set_text_if(slide, pred, new: str, all_matches: bool = False) -> bool:
    found = False
    for sh in slide.shapes:
        if not getattr(sh, "has_text_frame", False):
            continue
        if pred(sh.text_frame.text.strip()):
            _set_run_text(sh, new)
            found = True
            if not all_matches:
                return True
    return found


def _table_data(table) -> list[list[str]]:
    return [
        [table.cell(r, c).text_frame.text for c in range(len(table.columns))]
        for r in range(len(table.rows))
    ]


def _add_equip_table(slide, data: list[list[str]], l: float, t: float, w: float) -> None:
    n_rows = len(data)
    max_h = 5.36
    row_h = min(0.38, max_h / n_rows)
    table_h = row_h * n_rows
    font = 12 if n_rows > 16 else 13
    tbl_shape = slide.shapes.add_table(n_rows, 6, Inches(l), Inches(t), Inches(w), Inches(table_h))
    table = tbl_shape.table
    widths = [3.15, 2.15, 2.55, 1.35, 1.55, 1.683]
    for i, cw in enumerate(widths):
        table.columns[i].width = Inches(cw)
    for row in table.rows:
        row.height = Inches(row_h)
    aligns = ["left", "left", "left", "center", "right", "center"]
    for r, row in enumerate(data):
        if r == 0:
            for c, val in enumerate(row):
                fill_cell(table.cell(r, c), val, 12, True, WHITE, NAVY, aligns[c])
            continue
        bg = ROW_ALT if r % 2 == 0 else WHITE
        ne = row[3]
        for c, val in enumerate(row):
            col = MUTED if c == 3 and ne == "既有" else (HIPO_C if c == 3 else TEXT)
            fill_cell(table.cell(r, c), val, font, c in (3, 4, 5), col, bg, aligns[c])


def _replace_table(slide, data: list[list[str]]) -> None:
    tables = [sh for sh in slide.shapes if getattr(sh, "has_table", False)]
    if not tables:
        raise RuntimeError("table not found")
    sh = tables[0]
    l, t, w = _emu_in(sh.left), _emu_in(sh.top), _emu_in(sh.width)
    _remove(sh)
    _add_equip_table(slide, data, l, t, w)


def _renumber(prs) -> None:
    total = len(prs.slides)
    pat = re.compile(r"^\d+\s*/\s*\d+$")
    for i, slide in enumerate(prs.slides, 1):
        _set_text_if(
            slide,
            lambda t, p=pat: bool(p.match(t)),
            f"{i} / {total}",
            all_matches=True,
        )


def patch(src: Path = V1, dests: list[Path] | None = None) -> Presentation:
    dests = dests or [V1, OUT]
    prs = Presentation(str(src))
    for idx, key in GANTT_SLIDES.items():
        replace_gantt(prs.slides[idx], GANTT_SPECS[key])

    crd_idx = 31
    crd_slide = prs.slides[crd_idx]
    tables = [sh for sh in crd_slide.shapes if getattr(sh, "has_table", False)]
    data = _table_data(tables[0].table)
    header, rows = data[0], data[1:]
    mid = 12
    page1 = [header] + rows[:mid]
    page2 = [header] + rows[mid:]

    _clone_slide(prs, crd_idx)
    _move_slide(prs, len(prs.slides) - 1, crd_idx + 1)

    _replace_table(prs.slides[crd_idx], page1)
    if not _set_text_if(prs.slides[crd_idx], lambda t: t == "CRD", "CRD（1/2）"):
        raise RuntimeError("CRD subtitle not found on page 1")

    _replace_table(prs.slides[crd_idx + 1], page2)
    if not _set_text_if(prs.slides[crd_idx + 1], lambda t: t in ("CRD", "CRD（1/2）"), "CRD（2/2）"):
        raise RuntimeError("CRD subtitle not found on page 2")

    _renumber(prs)
    for dest in dests:
        dest.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(dest))
        print("wrote", dest, "slides", len(prs.slides), "bytes", dest.stat().st_size)
    return prs


if __name__ == "__main__":
    patch()
