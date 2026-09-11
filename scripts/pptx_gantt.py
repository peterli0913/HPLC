#!/usr/bin/env python3
"""Native PowerPoint Gantt charts for the 9.7 portfolio briefing.

Dates and class colours follow the live HTML. The extension window opens on
2026-05-01 so the completed RIBA 1 bar is not clipped.
"""
from __future__ import annotations

import calendar
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_hplc_lyopho_briefing import GANTT_CALENDAR  # noqa: E402
from build_management_briefing import GANTT_JS  # noqa: E402
from c1_oeb5_cost import GANTT_C1  # noqa: E402
from hipo_lab_cost import GANTT_HIPO, SHIFT_WEEKS_CLARE  # noqa: E402

FONT = "微软雅黑"
TODAY = datetime(2026, 8, 21)
SHIFT_MONTHS = 3
HIPO_SHIFT_DAYS = SHIFT_WEEKS_CLARE * 7
HIPO_DELIVERY = "2027-11-12"

NAVY = RGBColor(0x0F, 0x2B, 0x46)
TEAL = RGBColor(0x00, 0x96, 0x88)
TEXT = RGBColor(0x2C, 0x3E, 0x50)
MUTED = RGBColor(0x5A, 0x6A, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE8, 0xEE)
TRACK = RGBColor(0xEE, 0xF2, 0xF6)
ZEBRA = RGBColor(0xF4, 0xF7, 0xFA)
GRID = RGBColor(0xE8, 0xEE, 0xF4)
AXIS = RGBColor(0xDD, 0xE3, 0xE8)
TODAY_C = RGBColor(0xB4, 0x3A, 0x2A)
CARD = RGBColor(0xFF, 0xFF, 0xFF)

CLS = {
    "done": RGBColor(0x00, 0x96, 0x88),
    "plan": RGBColor(0x1A, 0x4A, 0x6E),
    "build": RGBColor(0x2E, 0x6D, 0xA4),
    "warn": RGBColor(0xC9, 0xA2, 0x27),
    "assume": RGBColor(0x6D, 0x5B, 0x95),
    "staff": RGBColor(0x5B, 0x6E, 0xAE),
}

EXT_LABELS = {
    "gFs": "可行性 RIBA 1",
    "gR2": "概念 RIBA 2",
    "gR3": "方案 RIBA 3",
    "gPlan": "规划",
    "gR4": "详细 RIBA 4",
    "gEquip": "长周期设备",
    "gMed": "中等周期设备",
    "gPre": "施工准备",
    "gR5": "施工 RIBA 5",
    "gComm": "调试",
    "gEnd": "验证/竣工",
}
HPLC_LABELS = {
    "gFs": "FS / 基准",
    "gEng": "Engineer 就位",
    "gFeed": "FEED（示意）",
    "gDd": "详细设计（示意）",
    "gLySpec": "冻干规格 / 资金",
    "gLyMfg": "冻干制造",
    "gLyFat": "冻干 FAT",
    "gLyShip": "冻干运输安装",
    "gLyVal": "冻干验证 → PQ",
    "gHplcSpec": "HPLC 规格 / 资金",
    "gHplcMfg": "HPLC 制造",
    "gHplcFat": "HPLC FAT",
    "gHplcShip": "HPLC 运输安装",
    "gTanks": "移动头罐",
    "gWaste": "废液罐",
    "gRetrofit": "改造",
}
C1_LABELS = {
    "c1Scope": "范围定稿",
    "c1Dd": "详细设计",
    "c1Build": "下单及设备制造",
    "c1IQ": "安装与确认",
}
HIPO_LABELS = {
    "hFund": "资金批准节点",
    "hConsult": "顾问任命 + BoD",
    "hContractor": "主承包商招标任命",
    "hAward": "合同授予与动员",
    "hSurvey": "勘查与项目控制",
    "hLabDesign": "实验室设计",
    "hCdm": "CDM",
    "hIso": "隔离器采购制造",
    "hFume": "通风柜采购（待定）",
    "hFurn": "实验室家具采购",
    "hTrade": "分包工程招标",
    "hConstr": "施工",
    "hComm": "调试与移交",
    "hRisk": "进度风险预留 → 交付",
}


def _d(s: str) -> datetime:
    return datetime.strptime(s, "%Y-%m-%d")


def _add_months(date_str: str, months: int) -> str:
    d = datetime.strptime(date_str, "%Y-%m-%d")
    total = d.year * 12 + (d.month - 1) + months
    year, month = divmod(total, 12)
    month += 1
    day = min(d.day, calendar.monthrange(year, month)[1])
    return f"{year:04d}-{month:02d}-{day:02d}"


def _shift_label(label: str, months: int) -> str:
    def repl(m: re.Match) -> str:
        total = int(m.group(1)) * 12 + int(m.group(2)) - 1 + months
        year, month = divmod(total, 12)
        return f"{year:04d}/{month + 1:02d}"

    return re.sub(r"(\d{4})/(\d{2})", repl, label)


def shift_gantt(rows: list, months: int = SHIFT_MONTHS) -> list:
    shifted = []
    for row in rows:
        r = list(row)
        if r[3] == "done":
            shifted.append(r)
            continue
        r[1] = _add_months(r[1], months)
        r[2] = _add_months(r[2], months)
        if len(r) > 6 and isinstance(r[6], str):
            r[6] = _shift_label(r[6], months)
        if len(r) > 7 and isinstance(r[7], str):
            r[7] = _add_months(r[7], months)
        shifted.append(r)
    return shifted


def shift_gantt_days(rows: list, days: int) -> list:
    out = []
    for row in rows:
        r = list(row)
        r[1] = (_d(r[1]) + timedelta(days=days)).strftime("%Y-%m-%d")
        r[2] = (_d(r[2]) + timedelta(days=days)).strftime("%Y-%m-%d")
        out.append(r)
    return out


def date_label(row: list) -> str:
    if len(row) > 6 and isinstance(row[6], str) and "–" in row[6]:
        return row[6]
    start = (row[6] if len(row) > 6 and isinstance(row[6], str) else row[1])[:7].replace("-", "/")
    end = (row[7] if len(row) > 7 and isinstance(row[7], str) else row[2])[:7].replace("-", "/")
    return f"{start} – {end}"


def bars_from(rows: list, labels: dict[str, str]) -> list[dict]:
    out = []
    for row in rows:
        out.append(
            {
                "id": row[0],
                "label": labels[row[0]],
                "start": _d(row[1]),
                "end": _d(row[2]),
                "cls": row[3],
                "dates": date_label(row),
            }
        )
    return out


def _ticks(*pairs: tuple[str, str]) -> list[tuple[datetime, str]]:
    return [(_d(d), lab) for d, lab in pairs]


GANTT_HPLC_SHIFTED = shift_gantt(GANTT_CALENDAR)
for _row in GANTT_HPLC_SHIFTED:
    if _row[0] == "gLyVal":
        _row[2] = "2028-06-08"
GANTT_C1_SHIFTED = shift_gantt(GANTT_C1)
GANTT_HIPO_SHIFTED = shift_gantt_days(GANTT_HIPO, HIPO_SHIFT_DAYS)
GANTT_HIPO_SHIFTED[-1][2] = HIPO_DELIVERY

GANTT_SPECS = {
    "ext": {
        "bars": bars_from(GANTT_JS, EXT_LABELS),
        "t0": _d("2026-05-01"),
        "t1": _d("2030-05-07"),
        "ticks": _ticks(
            ("2026-05-01", "2026"),
            ("2027-01-01", "2027"),
            ("2028-01-01", "2028"),
            ("2029-01-01", "2029"),
            ("2030-01-01", "2030"),
        ),
        "legend": [
            ("done", "已完成"),
            ("plan", "设计"),
            ("build", "施工"),
            ("warn", "节点"),
        ],
        "today": TODAY,
    },
    "hplc": {
        "bars": bars_from(GANTT_HPLC_SHIFTED, HPLC_LABELS),
        "t0": _d("2026-05-01"),
        "t1": _d("2028-06-30"),
        "ticks": _ticks(
            ("2026-05-01", "2026 H1"),
            ("2026-07-01", "2026 H2"),
            ("2027-01-01", "2027 H1"),
            ("2027-07-01", "2027 H2"),
            ("2028-01-01", "2028 H1"),
        ),
        "legend": [
            ("done", "完成"),
            ("plan", "采购 / 制造"),
            ("build", "施工 / 验证"),
            ("warn", "关键路径"),
            ("staff", "工程师就位"),
            ("assume", "§4.3 示意"),
        ],
        "today": TODAY,
    },
    "c1": {
        "bars": bars_from(GANTT_C1_SHIFTED, C1_LABELS),
        "t0": _d("2026-08-01"),
        "t1": _d("2027-10-31"),
        "ticks": _ticks(
            ("2026-08-01", "2026 H2"),
            ("2027-01-01", "2027 H1"),
            ("2027-07-01", "2027 H2"),
        ),
        "legend": [
            ("plan", "设计 / 准备"),
            ("warn", "下单 + 制造"),
            ("build", "施工 / 确认"),
        ],
        "today": TODAY,
    },
    "hipo": {
        "bars": bars_from(GANTT_HIPO_SHIFTED, HIPO_LABELS),
        "t0": _d("2026-09-01"),
        "t1": _d("2027-12-15"),
        "ticks": _ticks(
            ("2026-09-01", "2026 H2"),
            ("2027-01-01", "2027 H1"),
            ("2027-07-01", "2027 H2"),
        ),
        "legend": [
            ("plan", "设计 / 合约"),
            ("warn", "采购 / 制造 / 预留"),
            ("build", "施工 / 调试"),
        ],
        "today": TODAY,
    },
}


def _set_run_font(run, size: int, bold: bool, color: RGBColor):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def _no_line(shape):
    shape.line.fill.background()


def _rect(slide, l, t, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    if w <= 0 or h <= 0:
        return None
    s = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    _no_line(s)
    return s


def _round(slide, l, t, w, h, fill, adj=0.18):
    s = _rect(slide, l, t, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE)
    if s is None:
        return None
    try:
        s.adjustments[0] = adj
    except Exception:
        pass
    return s


def _tb(slide, l, t, w, h, text, size, bold, color, align=PP_ALIGN.LEFT, anchor="ctr"):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    body = tf._txBody.find(qn("a:bodyPr"))
    if body is not None:
        body.set("anchor", anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size, bold, color)
    return box


def _frac(d: datetime, t0: datetime, t1: datetime) -> float:
    span = (t1 - t0).total_seconds()
    if span <= 0:
        return 0.0
    return max(0.0, min(1.0, (d - t0).total_seconds() / span))


def _layout(n: int, height: float) -> dict:
    axis_h = 0.30
    legend_h = 0.34
    pad = 0.10
    avail = max(1.2, height - axis_h - legend_h - pad)
    if n <= 4:
        cap = 1.08
        name_w, date_w, lab_pt, date_pt = 2.55, 1.72, 14, 12
    elif n <= 8:
        cap = 0.62
        name_w, date_w, lab_pt, date_pt = 2.48, 1.68, 13, 11
    elif n <= 12:
        cap = 0.50
        name_w, date_w, lab_pt, date_pt = 2.42, 1.65, 12, 10
    else:
        cap = 0.36
        name_w, date_w, lab_pt, date_pt = 2.38, 1.62, 10, 9
    row_h = min(cap, avail / n)
    block_h = row_h * n
    extra = max(0.0, avail - block_h)
    return {
        "axis_h": axis_h,
        "legend_h": legend_h,
        "pad": pad,
        "name_w": name_w,
        "date_w": date_w,
        "lab_pt": lab_pt,
        "date_pt": date_pt,
        "row_h": row_h,
        "block_h": block_h,
        "y0": axis_h + extra * 0.22,
    }


def draw_gantt(slide, l: float, t: float, w: float, h: float, spec: dict):
    bars = spec["bars"]
    t0, t1 = spec["t0"], spec["t1"]
    ticks = spec["ticks"]
    legend = spec["legend"]
    today = spec.get("today")
    n = len(bars)
    ly = _layout(n, h)

    _round(slide, l, t, w, h, CARD, 0.04)
    rim = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    rim.fill.background()
    rim.line.color.rgb = LINE
    rim.line.width = Pt(1)
    try:
        rim.adjustments[0] = 0.04
    except Exception:
        pass

    name_w = ly["name_w"]
    date_w = ly["date_w"]
    track_l = l + name_w
    track_w = w - name_w - date_w - 0.10
    axis_t = t + 0.06
    rows_t = t + ly["y0"]
    row_h = ly["row_h"]

    _rect(slide, track_l, axis_t + ly["axis_h"] - 0.03, track_w, 0.012, AXIS)

    for dt, lab in ticks:
        x = track_l + track_w * _frac(dt, t0, t1)
        _rect(slide, x, rows_t, 0.012, ly["block_h"], GRID)
        tw = 0.90 if len(lab) > 4 else 0.70
        align = PP_ALIGN.LEFT if dt <= t0 else (PP_ALIGN.RIGHT if _frac(dt, t0, t1) > 0.92 else PP_ALIGN.CENTER)
        _tb(slide, x - tw / 2, axis_t, tw, 0.24, lab, 10, True, MUTED, align)

    for i, b in enumerate(bars):
        y = rows_t + i * row_h
        if i % 2:
            _rect(slide, l + 0.08, y, w - 0.16, row_h, ZEBRA)
        _tb(
            slide,
            l + 0.10,
            y,
            name_w - 0.16,
            row_h,
            b["label"],
            ly["lab_pt"],
            True,
            NAVY,
            PP_ALIGN.RIGHT,
        )
        track_h = min(0.62, max(0.18, row_h * 0.56))
        track_y = y + (row_h - track_h) / 2
        _round(slide, track_l, track_y, track_w, track_h, TRACK, 0.35)
        x0 = track_l + track_w * _frac(b["start"], t0, t1)
        x1 = track_l + track_w * _frac(b["end"], t0, t1)
        days = max(1, (b["end"] - b["start"]).days)
        min_w = 0.16 if days <= 14 else 0.12
        bar_w = max(x1 - x0, min_w)
        if x0 + bar_w > track_l + track_w:
            bar_w = max(0.10, track_l + track_w - x0)
        bar_h = max(0.14, track_h - 0.04)
        bar_y = track_y + (track_h - bar_h) / 2
        _round(slide, x0, bar_y, bar_w, bar_h, CLS[b["cls"]], 0.42)
        _tb(
            slide,
            track_l + track_w + 0.08,
            y,
            date_w - 0.06,
            row_h,
            b["dates"],
            ly["date_pt"],
            False,
            MUTED,
            PP_ALIGN.LEFT,
        )

    if today is not None and t0 <= today <= t1:
        tx = track_l + track_w * _frac(today, t0, t1)
        _rect(slide, tx, rows_t, 0.018, ly["block_h"], TODAY_C)
        if _frac(today, t0, t1) < 0.10:
            _tb(slide, tx + 0.05, axis_t, 0.48, 0.22, "约今", 9, True, TODAY_C, PP_ALIGN.LEFT)
        else:
            _tb(slide, tx - 0.24, axis_t, 0.48, 0.22, "约今", 9, True, TODAY_C, PP_ALIGN.CENTER)

    _draw_legend(slide, legend, l + 0.20, t + h - ly["legend_h"] + 0.02, w - 0.40, ly["legend_h"] - 0.04)


def _draw_legend(slide, items, l, t, w, h):
    n = len(items)
    if n == 0:
        return
    widths = [0.22 + max(0.72, 0.155 * len(lab)) for _, lab in items]
    gap = 0.36
    total = sum(widths) + gap * (n - 1)
    x = l + max(0.0, (w - total) / 2)
    for (cls, lab), unit in zip(items, widths):
        _round(slide, x, t + (h - 0.12) / 2, 0.12, 0.12, CLS[cls], 0.35)
        _tb(slide, x + 0.18, t, unit - 0.18, h, lab, 10, False, MUTED, PP_ALIGN.LEFT)
        x += unit + gap
